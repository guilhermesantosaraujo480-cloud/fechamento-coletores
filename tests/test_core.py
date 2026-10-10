import unittest
from datetime import date
from io import BytesIO

import tests  # noqa: F401  (instala o stub do bcrypt, se preciso)
from core import armazenamento as arm
from core import auth, consulta
from core import exportacao as exp
from core import financeiro as fin
from core import utils as u
from tests.fakes import FakeSupabase


class TestUtils(unittest.TestCase):
    def test_brl(self):
        self.assertEqual(u.brl(1234.5), "R$ 1.234,50")
        self.assertEqual(u.brl(-3), "-R$ 3,00")
        self.assertEqual(u.brl(None), "R$ 0,00")
        self.assertEqual(u.brl(float("nan")), "R$ 0,00")

    def test_datas_e_periodos(self):
        self.assertEqual(u.data_br("2026-10-09T10:00:00"), "09/10/2026")
        self.assertEqual(u.data_br("lixo"), "lixo")
        self.assertEqual(u.periodo_por_nome("Mês passado", date(2026, 1, 10)), (date(2025, 12, 1), date(2025, 12, 31)))
        self.assertEqual(u.periodo_por_nome("Últimos 7 dias", date(2026, 10, 10))[0], date(2026, 10, 4))
        self.assertEqual(u.dia_seguinte_iso(date(2026, 12, 31)), "2027-01-01")

    def test_centavos_e_busca(self):
        self.assertEqual(u.para_centavos(0.1) + u.para_centavos(0.2), 30)
        self.assertEqual(u.para_centavos("x"), 0)
        self.assertEqual(u.limpar_termo_busca("a%b,(c)_'"), "a b  c")


class TestFinanceiro(unittest.TestCase):
    def setUp(self):
        self.c = fin.df_coletas([
            {"id": 1, "data": "2026-10-01", "coletor": "Ana", "quantidade": 3, "status": "Aprovado", "valor_total": 30.0, "pago": False},
            {"id": 2, "data": "2026-10-02T10:00", "coletor": "Ana", "quantidade": 2, "status": "Pendente", "valor_total": 20.0, "pago": None},
            {"id": 3, "data": "2026-10-02", "coletor": "Bia", "quantidade": 1, "status": "Aprovado", "valor_total": 0.1, "pago": True},
        ])
        self.v = fin.df_vales([{"id": 1, "data": "2026-10-03", "coletor": "Ana", "valor_vale": 5.5}])
        self.p = fin.df_premios([{"id": 1, "data": "2026-10-03", "coletor": "Ana", "valor_premiacao": 0.2}])

    def test_fechamento_so_conta_aprovadas(self):
        f = fin.calcular_fechamento(self.c, self.v, self.p)
        self.assertEqual((f.aparelhos, f.bruto, f.premios, f.vales, f.liquido), (4, 3010, 20, 550, 2480))
        self.assertEqual((f.nao_pago, f.pendentes), (3000, 1))

    def test_vazio(self):
        f = fin.calcular_fechamento(fin.df_coletas([]), fin.df_vales([]), fin.df_premios([]))
        self.assertEqual(f.liquido, 0)

    def test_resumo_total_e_serie(self):
        r = fin.adicionar_total(fin.resumo_por_coletor(self.c, self.v, self.p, ["Caio"]))
        self.assertEqual(list(r["Coletor"]), ["Ana", "Bia", "Caio", "TOTAL GERAL"])
        self.assertAlmostEqual(r.iloc[-1]["Líquido"], 24.8)
        s = fin.serie_diaria(self.c, date(2026, 10, 1), date(2026, 10, 4))
        self.assertEqual(list(s["Aparelhos"]), [3, 1, 0, 0])
        longa = fin.serie_diaria(self.c, date(2026, 1, 1), date(2026, 10, 4))
        self.assertLess(len(longa), 12)

    def test_periodo_com_timestamp(self):
        self.assertEqual(len(fin.filtrar_periodo(self.c, date(2026, 10, 2), date(2026, 10, 2))), 2)


class TestConsulta(unittest.TestCase):
    def test_paginacao_passa_de_1000(self):
        db = FakeSupabase(coletas=[{"id": i, "data": "2026-10-01", "coletor": "Ana"} for i in range(1, 2501)])
        self.assertEqual(len(consulta.buscar_todos(db, "coletas")), 2500)

    def test_faixa_inclui_o_dia_final_mesmo_com_hora(self):
        db = FakeSupabase(coletas=[{"id": 1, "data": "2026-10-05T23:59:00"}, {"id": 2, "data": "2026-10-06"},
                                   {"id": 3, "data": "2026-10-04"}])
        ids = [r["id"] for r in consulta.buscar_todos(db, "coletas", faixa=("data", date(2026, 10, 4), date(2026, 10, 5)))]
        self.assertEqual(sorted(ids), [1, 3])

    def test_in_e_contar(self):
        db = FakeSupabase(coletas=[{"id": i, "status": "Pendente" if i % 2 else "Aprovado"} for i in range(1, 11)])
        self.assertEqual(len(consulta.buscar_todos(db, "coletas", filtros={"id": [1, 2, 3]})), 3)
        self.assertEqual(consulta.contar(db, "coletas", {"status": "Pendente"}), 5)


class TestAuth(unittest.TestCase):
    def novo_db(self):
        return FakeSupabase(usuarios=[
            {"id": 1, "usuario": "ana", "senha": auth.gerar_hash_senha("segredo1"), "nome_completo": "Ana", "cargo": "COLETOR"},
            {"id": 2, "usuario": "velho", "senha": "texto123", "nome_completo": "Velho", "cargo": "ADM"},
            {"id": 3, "usuario": "off", "senha": auth.gerar_hash_senha("segredo1"), "nome_completo": "Off", "cargo": "COLETOR", "ativo": False},
        ])

    def test_login_ok_errado_e_inativo(self):
        db, lim = self.novo_db(), auth.Limitador()
        r = auth.autenticar(db, " ANA ", "segredo1", lim)
        self.assertTrue(r.ok)
        self.assertNotIn("senha", r.usuario)
        self.assertNotEqual(db.tabelas["usuarios"][0]["session_token"], r.token)  # só o hash é guardado
        self.assertEqual(auth.usuario_por_token(db, r.token)["usuario"], "ana")
        self.assertFalse(auth.autenticar(db, "ana", "errada", lim).ok)
        self.assertIn("desativado", auth.autenticar(db, "off", "segredo1", lim).mensagem)

    def test_migra_senha_em_texto_puro(self):
        db = self.novo_db()
        self.assertTrue(auth.autenticar(db, "velho", "texto123", auth.Limitador()).ok)
        self.assertTrue(db.tabelas["usuarios"][1]["senha"].startswith("$2"))

    def test_bloqueio_apos_5_falhas(self):
        db, t = self.novo_db(), [1000.0]
        lim = auth.Limitador(relogio=lambda: t[0])
        for _ in range(5):
            auth.autenticar(db, "ana", "x", lim)
        self.assertIn("Muitas tentativas", auth.autenticar(db, "ana", "segredo1", lim).mensagem)
        t[0] += 301
        self.assertTrue(auth.autenticar(db, "ana", "segredo1", lim).ok)

    def test_trocar_senha(self):
        db = self.novo_db()
        self.assertFalse(auth.alterar_propria_senha(db, 1, "errada", "novasenha")[0])
        self.assertFalse(auth.alterar_propria_senha(db, 1, "segredo1", "123")[0])
        self.assertTrue(auth.alterar_propria_senha(db, 1, "segredo1", "novasenha")[0])
        self.assertTrue(auth.autenticar(db, "ana", "novasenha", auth.Limitador()).ok)


class TestArmazenamento(unittest.TestCase):
    def test_imagem_invalida_e_exif(self):
        from PIL import Image
        with self.assertRaises(ValueError):
            arm.processar_imagem(b"isso nao e imagem")
        img = Image.new("RGB", (400, 200), "red")
        exif = Image.Exif()
        exif[0x0112] = 6  # girar 90°
        buf = BytesIO()
        img.save(buf, "JPEG", exif=exif)
        saida = Image.open(BytesIO(arm.processar_imagem(buf.getvalue())))
        self.assertEqual(saida.size, (200, 400))

    def test_png_transparente_e_upload(self):
        from PIL import Image
        buf = BytesIO()
        Image.new("RGBA", (50, 50), (0, 0, 0, 0)).save(buf, "PNG")
        db = FakeSupabase()
        caminho = arm.salvar_foto(db, buf.getvalue(), "coleta_ana/../x")
        self.assertIn(caminho, db.arquivos)
        self.assertNotIn("/", caminho)

    def test_referencias(self):
        self.assertEqual(arm.referencia_foto(None), (None, None))
        self.assertEqual(arm.referencia_foto("nan"), (None, None))
        self.assertEqual(arm.referencia_foto("x/y.jpg"), ("caminho", "x/y.jpg"))
        antiga = "https://p.supabase.co/storage/v1/object/public/comprovantes/a%20b.jpg?t=1"
        self.assertEqual(arm.referencia_foto(antiga), ("caminho", "a b.jpg"))
        self.assertEqual(arm.referencia_foto("https://outro.com/a.jpg"), ("url", "https://outro.com/a.jpg"))


class TestExportacao(unittest.TestCase):
    def test_formulas_e_excel(self):
        import pandas as pd
        from openpyxl import load_workbook
        df = pd.DataFrame({"a": ["=1+1", "-R$ 10,00", "-cmd", "normal"]})
        self.assertEqual(list(exp.neutralizar_formulas(df)["a"]), ["'=1+1", "-R$ 10,00", "'-cmd", "normal"])
        wb = load_workbook(BytesIO(exp.gerar_excel({"Resumo: [x]": df})))
        self.assertEqual(wb.sheetnames, ["Resumo- -x-"])


if __name__ == "__main__":
    unittest.main()
