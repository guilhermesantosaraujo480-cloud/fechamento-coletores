import io
import unittest

import pandas as pd

from tests import fakes

ST = fakes.instalar()

from rotas import envio, leitura, mensagens, perfis as pf, processamento as pr, vinculos  # noqa: E402
from rotas import whatsapp as wa  # noqa: E402


def df_reintegra():
    return pd.DataFrame({
        "BA": ["1", "2", "3", "4"], "Cliente": ["Ana", "Beto", "Caio", "Dani"],
        "Endereço": ["Rua A", "Rua B", "Rua C", "Rua A"], "Número": ["10.0", "5", "", "10"],
        "Bairro": ["Centro", "Centro", "Lapa", "Centro"], "Cep": ["1310000", "01310-001", "x", "1310000"],
        "Cidade": ["SP"] * 4, "UF": ["SP"] * 4, "Telefone 1": ["11987654321", "5511987654322", "", "1133334444"],
        "Complemento": [""] * 4})


def df_tim():
    return pd.DataFrame({
        "Número do Pedido": ["9", "8"], "Conta.2": ["c1", "c2"], "Nome": ["X", "Y"], "Logradouro": ["Rua Z", "Rua W"],
        "Número": ["1", "2"], "Bairro": ["B1", "B2"], "CEP": ["01000000", "02000000"], "Cidade": ["SP", "SP"],
        "Telefone": ["", "11911112222"], "Celular": ["11922223333", ""], "Informação adicional": ["", ""], "UF": ["SP", "SP"]})


class TestPerfis(unittest.TestCase):
    def test_deteccao(self):
        self.assertEqual(pf.detectar_perfil(list(df_reintegra().columns), pf.PERFIS_PADRAO).nome, pf.REINTEGRA.nome)
        self.assertEqual(pf.detectar_perfil(list(df_tim().columns), pf.PERFIS_PADRAO).nome, pf.TIM_CSS.nome)
        self.assertIsNone(pf.detectar_perfil(["a", "b"], pf.PERFIS_PADRAO))

    def test_dict_roundtrip(self):
        p = pf.Perfil.de_dict(pf.TIM_CSS.para_dict())
        self.assertEqual(p.mapa, pf.TIM_CSS.mapa)


class TestProcessamento(unittest.TestCase):
    def test_telefone_cep(self):
        self.assertEqual(pr.normalizar_telefone("5511987654321"), ("(11) 98765-4321", True))
        self.assertEqual(pr.normalizar_telefone("1133334444"), ("(11) 3333-4444", True))
        self.assertFalse(pr.normalizar_telefone("123")[1])
        self.assertEqual(pr.normalizar_cep("1310000"), ("01310-000", True))
        self.assertFalse(pr.normalizar_cep("x")[1])
        self.assertEqual(pr.numero_whatsapp("(11) 98765-4321"), "5511987654321")
        self.assertEqual(pr.numero_whatsapp(""), "")
        self.assertEqual(pr.numero_whatsapp("5511987654321"), "5511987654321")

    def test_reintegra(self):
        can, falt = pr.aplicar_perfil(df_reintegra(), pf.REINTEGRA)
        self.assertEqual(falt, [])
        self.assertEqual(can["endereco_completo"].iloc[0], "Rua A N° 10")
        self.assertEqual(can["endereco_completo"].iloc[2], "Rua C")
        probs = pr.checar_qualidade(can)
        self.assertIn("Sem telefone válido", set(probs["Problema"]))
        self.assertIn("CEP inválido", set(probs["Problema"]))
        saida = pr.montar_saida(can, pf.REINTEGRA)
        self.assertEqual(list(saida.columns)[:3], [c for c, _ in pf.REINTEGRA.saida][:3])
        self.assertIn("Mapa", saida.columns)

    def test_tim_e_celular_vale(self):
        can, falt = pr.aplicar_perfil(df_tim(), pf.TIM_CSS)
        self.assertEqual(falt, [])
        self.assertTrue(can["tel_ok"].all())  # telefone OU celular

    def test_coluna_faltando(self):
        _, falt = pr.aplicar_perfil(df_reintegra().drop(columns=["Bairro"]), pf.REINTEGRA)
        self.assertEqual(falt, ["Bairro"])

    def test_divisao_equilibrada(self):
        n = 60
        df = pd.DataFrame({"BA": range(n), "Cliente": ["c"] * n, "Endereço": ["r"] * n, "Número": ["1"] * n,
                           "Bairro": [f"B{i % 12}" for i in range(n)], "Cep": ["01000000"] * n, "Cidade": ["SP"] * n,
                           "UF": ["SP"] * n, "Telefone 1": ["11987654321"] * n, "Complemento": [""] * n})
        can, _ = pr.aplicar_perfil(df, pf.REINTEGRA)
        atrib = pr.dividir_entre(can, ["A", "B", "C"])
        cont = atrib.value_counts()
        self.assertEqual(set(cont.index), {"A", "B", "C"})
        self.assertLessEqual(cont.max() - cont.min(), 10)
        self.assertEqual(len(atrib), n)
        # bairro gigante é repartido
        can["bairro"] = "UNICO"
        self.assertEqual(pr.dividir_entre(can, ["A", "B"]).nunique(), 2)
        self.assertTrue((pr.dividir_entre(can, []) == pr.SEM_COLETOR).all())

    def test_leitura_tab_latin1_e_excel(self):
        txt = "BA\tCliente\tEndereço\nx\tJosé\tRua\n".encode("latin1")
        class Arq(io.BytesIO):
            name = "a.xls"
        df = leitura.carregar_planilha(Arq(txt))
        self.assertEqual(df["Cliente"].iloc[0], "José")
        buf = Arq()
        buf.name = "a.xlsx"
        df_tim().to_excel(buf, index=False)
        buf.seek(0)
        self.assertEqual(leitura.carregar_planilha(buf)["Conta.2"].iloc[1], "c2")


class TestVinculos(unittest.TestCase):
    def test_sugestao(self):
        rota = [{"id": 1, "nome_rota": "GABRIEL", "usuario_id": None}, {"id": 2, "nome_rota": "RENILTON", "usuario_id": None},
                {"id": 3, "nome_rota": "CARLOS", "usuario_id": 9}]
        us = [{"id": 5, "usuario": "gabriel.s", "nome_completo": "Gabriel Souza", "cargo": "COLETOR"},
              {"id": 6, "usuario": "adm", "nome_completo": "Gabriel Chefe", "cargo": "ADM"},
              {"id": 9, "usuario": "carlos", "nome_completo": "Carlos", "cargo": "COLETOR"}]
        self.assertEqual([(c, u) for c, u, _ in vinculos.sugerir_vinculos(rota, us)], [(1, 5)])


class TestMensagem(unittest.TestCase):
    def test_limite(self):
        n = 300
        df = pd.DataFrame({"BA": range(n), "Cliente": ["c"] * n, "Endereço": ["r"] * n, "Número": ["1"] * n,
                           "Bairro": [f"Bairro numero {i}" for i in range(n)], "Cep": [""] * n, "Cidade": ["SP"] * n,
                           "UF": ["SP"] * n, "Telefone 1": [""] * n, "Complemento": [""] * n})
        can, _ = pr.aplicar_perfil(df, pf.REINTEGRA)
        from datetime import date
        self.assertLessEqual(len(mensagens.mensagem_rota("Reintegra", "GABRIEL", date(2026, 10, 10), can)), 900)
        self.assertLessEqual(len(mensagens.lista_para_copiar("Reintegra", "G", date(2026, 10, 10), can)), 3600)


class Resp:
    def __init__(self, status=200, corpo=None):
        self.status_code, self._c, self.text = status, corpo if corpo is not None else {}, ""

    def json(self):
        return self._c


class Http:
    def __init__(self, status=200, corpo=None):
        self.chamadas, self.status, self.corpo = [], status, corpo

    def _f(self, metodo):
        def f(url, **kw):
            self.chamadas.append((metodo, url, kw))
            return Resp(self.status, self.corpo)
        return f

    def __getattr__(self, m):
        return self._f(m)


class TestWhats(unittest.TestCase):
    def test_evolution(self):
        h = Http()
        g = wa.Evolution("https://e.test/", "inst", "KEY", http=h)
        g.enviar_texto("123@g.us", "oi")
        g.enviar_arquivo("123@g.us", b"abc", "r.xlsx", "leg")
        m, url, kw = h.chamadas[1]
        self.assertEqual(url, "https://e.test/message/sendMedia/inst")
        self.assertEqual(kw["headers"]["apikey"], "KEY")
        self.assertEqual(kw["json"]["fileName"], "r.xlsx")

    def test_zapi_e_erro(self):
        h = Http()
        wa.ZApi("https://z.test/instances/1/token/2", "CT", http=h).enviar_texto("5511", "x")
        self.assertEqual(h.chamadas[0][2]["headers"]["Client-Token"], "CT")
        with self.assertRaises(wa.ErroWhats):
            wa.Evolution("https://e.test", "i", "k", http=Http(401, {"e": "x"})).enviar_texto("1", "x")

    def test_criar_gateway(self):
        self.assertIsNone(wa.criar_gateway(lambda k, d="": d))
        cfg = {"WHATS_PROVEDOR": "zapi", "WHATS_API_URL": "https://z"}
        self.assertIsInstance(wa.criar_gateway(lambda k, d="": cfg.get(k, d)), wa.ZApi)

    def test_envio_aborta_apos_3_erros(self):
        class G:
            def __init__(self, falha):
                self.falha, self.n = falha, 0

            def enviar_arquivo(self, *a):
                self.n += 1
                if self.falha:
                    raise wa.ErroWhats("x")
        pacotes = [{"coletor": f"C{i}", "conteudo": b"", "arquivo": "a", "mensagem": "m", "telefone": ""} for i in range(5)]
        reg = []
        g = G(True)
        res, parada = envio.enviar_rotas(g, pacotes, ["g1"], lambda r, p: reg.append(r), dormir=lambda s: None)
        self.assertEqual(len(res), 3)
        self.assertTrue(parada)
        g = G(False)
        res, parada = envio.enviar_rotas(g, pacotes, ["g1", "g2"], lambda r, p: None, dormir=lambda s: None)
        self.assertEqual((len(res), parada), (10, ""))
        pacotes[0]["telefone"] = "5511"
        res, _ = envio.enviar_rotas(G(False), pacotes[:1], ["g1"], lambda r, p: None, individual=True, dormir=lambda s: None)
        self.assertEqual(len(res), 2)


if __name__ == "__main__":
    unittest.main()


# ------------------------------------------------------------------------- novidades
from datetime import date  # noqa: E402

from rotas import acompanhamento, avisos, distribuicao, nomes  # noqa: E402
from rotas import repositorio as repo  # noqa: E402

COLETORES = ["GEOVAN", "RAFAEL", "GABRIEL", "RENILTON", "ARLEIDE.JULIENE", "GUILHERME", "CARLOS", "DANILO",
             "CARLOS.ALBERTO", "JOSENILDA", "OUTROS"]


def df_reintegra_distribuido():
    """Relatório geral no estilo real: coluna Coletor com nomes completos, algumas OS sem dono."""
    linhas = []
    dados = [
        ("GUILHERME SANTOS ARAUJO", "Cliente ausente", "SAO PAULO", "BAIRRO A", "13/10/2026", "2"),
        ("RENILTON ARAUJO", "", "SAO PAULO", "BAIRRO A", "30/10/2026", "0"),
        ("Carlos Alberto dos Santos ", "", "ITAQUAQUECETUBA", "BAIRRO I", "", "0"),
        ("JULIENE ALVES MENDES DA SILVA ", "Cliente mudou-se", "ITAQUAQUECETUBA", "BAIRRO I", "", "0"),
        ("ARLEIDE SANTOS DA SIVA", "", "ITAQUAQUECETUBA", "BAIRRO I", "", "0"),
        ("", "", "SAO PAULO", "BAIRRO A", "14/10/2026", "0"),            # sem dono, bairro de Renilton
        ("", "Cliente mudou-se", "ITAQUAQUECETUBA", "BAIRRO I", "", "0"),  # sem dono, mudou-se
        ("", "Cliente ausente", "SAO PAULO", "BAIRRO Z", "", "1"),        # sem dono, já foi à rua
        ("", "", "ITAQUAQUECETUBA", "ITAQUAQUECETUBA", "", "0"),          # sem dono, bairro genérico
    ]
    for i, (col, motivo, cid, bairro, agenda, qtd) in enumerate(dados):
        linhas.append({"BA": f"BA{i}", "Cliente": f"Cli {i}", "Endereço": f"Rua {i}", "Número": str(i), "Bairro": bairro,
                       "Cep": f"0{i}000000", "Cidade": cid, "UF": "SP", "Telefone 1": "5511987654321", "Complemento": "",
                       "Coletor": col, "Ultimo Motivo/Ocorrencia": motivo, "Agenda": agenda, "Qtd Ag.": qtd,
                       "Tipo": "INVOLUNTARIO", "Backlog": "40", "Descrição Ocorrência": ""})
    return pd.DataFrame(linhas)


class TestNomes(unittest.TestCase):
    def test_casamento(self):
        casos = {"GUILHERME SANTOS ARAUJO": "GUILHERME", "Renilton Araujo da Cunha (Coletor)": "RENILTON",
                 "JULIENE ALVES MENDES DA SILVA ": "ARLEIDE.JULIENE", "ARLEIDE SANTOS DA SIVA": "ARLEIDE.JULIENE",
                 "GABRIEL SANTOS ARAÚJO": "GABRIEL", "CARLOS ALBERTO DOS SANTOS": "CARLOS.ALBERTO", "CARLOS": "CARLOS"}
        for bruto, esperado in casos.items():
            self.assertEqual(nomes.casar(bruto, COLETORES)[0], esperado, bruto)
        self.assertIsNone(nomes.casar("Fulano de Tal", COLETORES)[0])
        self.assertIsNone(nomes.casar("", COLETORES)[0])

    def test_apelido_vence_e_ajuste_da_tela(self):
        ap = {nomes.normalizar_nome("Fulano de Tal"): "DANILO"}
        self.assertEqual(nomes.casar("FULANO DE TAL", COLETORES, ap), ("DANILO", "apelido"))
        saida = nomes.atribuir(["Fulano de Tal", "", None, "Beltrano"], COLETORES, {}, {"Beltrano": "RAFAEL"})
        self.assertEqual(saida, [pr.SEM_COLETOR, pr.SEM_COLETOR, pr.SEM_COLETOR, "RAFAEL"])

    def test_ambiguo(self):
        self.assertEqual(nomes.casar("Ana Maria", ["ANA.PAULA", "ANA.CLARA"])[1], "ambíguo")


class TestAvisos(unittest.TestCase):
    def test_avisos_reintegra_e_css(self):
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        can = avisos.aplicar_avisos(can, date(2026, 10, 13))
        t = can["avisos"].tolist()
        self.assertIn("📅 Agendada para HOJE", t[0])
        self.assertIn("cliente ausente", t[0])
        self.assertTrue(can["urgente"].iloc[0])
        self.assertNotIn("Agend", t[1])            # 30/10 (data padrão) não vira aviso
        self.assertIn("amanhã", t[5])
        self.assertIn("mudou-se", t[3])
        self.assertIn("40 dias na fila", t[0])
        css = pd.DataFrame({"Número do Pedido": ["1"], "Conta.2": ["TIM | Vol"], "Nome": ["X"], "Logradouro": ["R"], "Número": ["1"],
                            "Bairro": ["B"], "CEP": ["01000000"], "Cidade": ["SP"], "Telefone": ["11911112222"], "Celular": [""],
                            "Informação adicional": [""], "UF": ["SP"], "Descrição da Prioridade": ["Zero - Extrema"],
                            "Descrição da Ocorrência": ["Área de Risco"], "Reaberto em": ["08/10/2026 11:46:41"],
                            "Transportador.1": ["Renilton Araujo da Cunha (Coletor)"]})
        c2, _ = pr.aplicar_perfil(css, pf.TIM_CSS)
        c2 = avisos.aplicar_avisos(c2, date(2026, 10, 13))
        self.assertIn("Prioridade máxima", c2["avisos"].iloc[0])
        self.assertIn("Área de risco", c2["avisos"].iloc[0])
        self.assertIn("Reaberta", c2["avisos"].iloc[0])
        self.assertEqual(c2["coletor"].iloc[0], "Renilton Araujo da Cunha (Coletor)")

    def test_planilha_sem_colunas_extras_nao_da_erro(self):
        can, falt = pr.aplicar_perfil(df_reintegra(), pf.REINTEGRA)  # sem Coletor/Agenda/Tipo...
        self.assertEqual(falt, [])
        can = avisos.aplicar_avisos(can, date(2026, 10, 13))
        self.assertTrue((can["coletor"] == "").all())
        self.assertFalse(can["urgente"].any())

    def test_urgentes_primeiro_e_saida_com_avisos(self):
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        can = avisos.aplicar_avisos(can, date(2026, 10, 13))
        ordenada = pr.ordenar_rota(can, urgentes_primeiro=True)
        self.assertTrue(ordenada["urgente"].iloc[0])
        saida = pr.montar_saida(ordenada, pf.REINTEGRA, incluir_mapa=True, incluir_avisos=True)
        self.assertEqual(list(saida.columns)[-2:], ["Mapa", "Avisos"])
        completa = pr.montar_saida(ordenada, pf.REINTEGRA, True, True, True)
        self.assertEqual(list(completa.columns)[-3:], ["Mapa", "WhatsApp", "Avisos"])
        self.assertEqual(list(pr.montar_saida(ordenada, pf.REINTEGRA, False, False).columns), [c for c, _ in pf.REINTEGRA.saida])

    def test_bairro_generico_agrupa_por_cep(self):
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        self.assertTrue(can["bairro_ef"].iloc[-1].startswith("CEP "))
        problemas = pr.checar_qualidade(can)
        self.assertTrue(problemas["Problema"].str.contains("Bairro genérico").any())

    def test_mensagem_traz_resumo_dos_avisos(self):
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        can = avisos.aplicar_avisos(can, date(2026, 10, 13))
        msg = mensagens.mensagem_rota("Reintegra", "GUILHERME", date(2026, 10, 13), can)
        self.assertIn("agendada(s) p/ hoje", msg)
        self.assertLessEqual(len(msg), 900)


class TestDistribuicao(unittest.TestCase):
    def test_sugestoes_com_regras(self):
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        atrib = pd.Series(nomes.atribuir(can["coletor"], COLETORES), index=can.index, dtype=object)
        tipos = {"GUILHERME": "retentativa", "ARLEIDE.JULIENE": "deposito"}
        s = distribuicao.sugerir(can, atrib, tipos)
        self.assertEqual(len(s), 4)
        por_ba = {can.at[i, "id"]: (r["sugestao"], r["razao"]) for i, r in s.iterrows()}
        self.assertEqual(por_ba["BA5"][0], "RENILTON")             # bairro dele
        self.assertEqual(por_ba["BA6"][0], "ARLEIDE.JULIENE")      # mudou-se -> depósito
        self.assertEqual(por_ba["BA7"][0], "GUILHERME")            # já foi à rua -> retentativas
        self.assertEqual(por_ba["BA8"][0], "CARLOS.ALBERTO")       # única que atende Itaquaquecetuba entre as normais
        # sem tipos especiais: tudo vai para coletores normais
        s2 = distribuicao.sugerir(can, atrib, {})
        self.assertNotIn(pr.SEM_COLETOR, set(s2["sugestao"]))

    def test_sem_ninguem_para_receber(self):
        can, _ = pr.aplicar_perfil(df_reintegra(), pf.REINTEGRA)
        atrib = pd.Series(pr.SEM_COLETOR, index=can.index, dtype=object)
        s = distribuicao.sugerir(can, atrib, {})
        self.assertTrue((s["sugestao"] == pr.SEM_COLETOR).all())

    def test_arquivo_para_plataforma(self):
        can, _ = pr.aplicar_perfil(df_reintegra(), pf.REINTEGRA)
        atrib = pd.Series(["GEOVAN", pr.SEM_COLETOR, "RAFAEL", "GEOVAN"], index=can.index, dtype=object)
        arq = distribuicao.arquivo_para_plataforma(can, atrib, pf.REINTEGRA)
        self.assertEqual(list(arq.columns), ["BA", "Coletor"])
        self.assertEqual(len(arq), 3)


class TestAcompanhamento(unittest.TestCase):
    def snap(self, ids_por_coletor, data):
        linhas = []
        for c, ids in ids_por_coletor.items():
            for i in ids:
                linhas.append({"chave": f"E|{data}|{i}", "empresa": "E", "data": data, "os_id": i, "coletor": c,
                               "bairro": "B", "cidade": "SP", "motivo": "", "urgente": False})
        return linhas

    def test_comparar(self):
        ant = self.snap({"A": ["1", "2", "3", "4"], "B": ["5", "6"]}, "2026-10-12")
        hoje = self.snap({"A": ["3", "4", "7"], "B": ["1", "6"]}, "2026-10-13")
        r = acompanhamento.comparar(ant, hoje)
        tab = r["por_coletor"].set_index("Coletor")
        self.assertEqual(tab.loc["A", "Saíram da lista"], 1)       # 2 sumiu; 1 foi para B
        self.assertEqual(tab.loc["A", "Passou para outro"], 1)
        self.assertEqual(tab.loc["B", "Recebeu de outro"], 1)
        self.assertEqual(tab.loc["B", "Saíram da lista"], 1)       # 5 sumiu
        self.assertEqual(tab.loc["A", "Entraram"], 1)
        self.assertEqual((r["total_antes"], r["total_hoje"]), (6, 5))
        prod = acompanhamento.produtividade({"2026-10-12": ant, "2026-10-13": hoje})
        self.assertEqual(int(prod.loc["2026-10-13", "A"]), 1)

    def test_snapshot_no_banco(self):
        sb = fakes.FakeSupabase()
        can, _ = pr.aplicar_perfil(df_reintegra_distribuido(), pf.REINTEGRA)
        can = avisos.aplicar_avisos(can, date(2026, 10, 13))
        atrib = pd.Series(nomes.atribuir(can["coletor"], COLETORES), index=can.index, dtype=object)
        linhas = acompanhamento.montar_snapshot(can, atrib, "Reintegra", date(2026, 10, 13))
        self.assertNotIn("nome", linhas[0])  # sem dados pessoais do cliente
        repo.salvar_snapshot(sb, linhas)
        repo.salvar_snapshot(sb, linhas)     # mesmo dia de novo: atualiza, não duplica
        self.assertEqual(len(sb.tabelas["rota_snapshots"]), len(linhas))
        linhas2 = acompanhamento.montar_snapshot(can, atrib, "Reintegra", date(2026, 10, 12))
        repo.salvar_snapshot(sb, linhas2)
        self.assertEqual(repo.datas_de_snapshot(sb, "Reintegra"), ["2026-10-13", "2026-10-12"])
        self.assertEqual(len(repo.ler_snapshot(sb, "Reintegra", "2026-10-13")), len(linhas))

    def test_apelidos_e_tipo_no_banco(self):
        sb = fakes.FakeSupabase(rota_coletores=[])
        repo.criar_coletor(sb, "novo", "11999990000", None, "retentativa")
        repo.criar_coletor(sb, "outro")
        lista = repo.listar_coletores(sb)
        self.assertEqual(repo.tipos_dos_coletores(lista), {"NOVO": "retentativa", "OUTRO": "normal"})
        self.assertNotIn("tipo", [c for c in lista if c["nome_rota"] == "OUTRO"][0])  # sem coluna nova quando é 'normal'
        repo.salvar_apelido(sb, "FULANO", "NOVO")
        repo.salvar_apelido(sb, "FULANO", "OUTRO")
        self.assertEqual(repo.listar_apelidos(sb), {"FULANO": "OUTRO"})
        repo.remover_apelido(sb, "FULANO")
        self.assertEqual(repo.listar_apelidos(sb), {})


class TestLinkWhatsapp(unittest.TestCase):
    def test_numero_e_link(self):
        self.assertEqual(pr.numero_do_celular("(11) 98765-4321"), "5511987654321")
        self.assertEqual(pr.numero_do_celular("5511987654321"), "5511987654321")
        self.assertEqual(pr.numero_do_celular("(11) 3333-4444"), "")      # fixo não tem WhatsApp
        self.assertEqual(pr.numero_do_celular(""), "")
        url = pr.link_whatsapp("(11) 98765-4321", "", "MARIA DA SILVA", "GEOVAN")
        self.assertTrue(url.startswith("https://wa.me/5511987654321?text="))
        from urllib.parse import unquote
        self.assertIn("Olá, Maria, tudo bem?", unquote(url))
        self.assertIn("retirar o equipamento", unquote(url))

    def test_usa_o_segundo_telefone_quando_o_primeiro_e_fixo(self):
        url = pr.link_whatsapp("(11) 3333-4444", "(11) 91111-2222", "Ana")
        self.assertIn("5511911112222", url)
        self.assertEqual(pr.link_whatsapp("(11) 3333-4444", "", "Ana"), "")

    def test_modelo_personalizado_sem_nome(self):
        from urllib.parse import unquote
        url = pr.link_whatsapp("11987654321", "", "", "RENILTON", "Olá, {cliente}! Sou {coletor}, cheguei.")
        self.assertIn("Olá! Sou RENILTON, cheguei.", unquote(url))
        self.assertIn("&", pr.link_whatsapp("11987654321", "", "X", "", "a & b") + "&")  # não quebra com caracteres especiais

    def test_coluna_na_planilha_e_clicavel(self):
        import io
        from openpyxl import load_workbook
        from core.exportacao import gerar_excel
        can, _ = pr.aplicar_perfil(df_reintegra(), pf.REINTEGRA)
        saida = pr.montar_saida(can, pf.REINTEGRA, True, False, True, None, "GEOVAN")
        self.assertTrue(saida["WhatsApp"].iloc[0].startswith("https://wa.me/"))
        self.assertEqual(saida["WhatsApp"].iloc[2], "")                    # linha sem telefone
        ws = load_workbook(io.BytesIO(gerar_excel({"Rota": saida})))["Rota"]
        cab = [c.value for c in ws[1]]
        col = cab.index("WhatsApp") + 1
        celula = ws.cell(row=2, column=col)
        self.assertEqual(celula.value, celula.hyperlink.target)           # texto = endereço (importa em roteirizador)
        self.assertIsNone(ws.cell(row=4, column=col).hyperlink)
